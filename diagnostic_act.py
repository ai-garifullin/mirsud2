import streamlit as st
import io
import os
import json
import re
from PIL import Image
import docx
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.styles.style import ParagraphStyle

DB_PATH = os.path.join(os.path.dirname(__file__), "courts_db.json")
BACKUP_CSV_PATH = r"E:\Базы\База Мировых.csv"


@st.cache_data
def load_courts_database():
    """Загружает базу данных судебных участков РТ из JSON или CSV."""
    if os.path.exists(DB_PATH):
        try:
            with open(DB_PATH, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass

    # Резервный источник: База Мировых.csv
    records = []
    csv_candidates = [BACKUP_CSV_PATH, "База Мировых.csv", "baza_mirovih.csv"]
    for cand in csv_candidates:
        if os.path.exists(cand):
            try:
                import csv
                with open(cand, "r", encoding="cp866", errors="replace") as f:
                    reader = csv.reader(f, delimiter=";")
                    next(reader, None)  # Пропуск заголовка
                    for row in reader:
                        if not row or not any(row):
                            continue
                        court_info = row[0].strip()
                        phone = row[1].strip() if len(row) > 1 else ""
                        email = row[2].strip() if len(row) > 2 else ""

                        m = re.search(r'(Судебный участок\s*\d+.*)', court_info, flags=re.IGNORECASE)
                        if m:
                            court_name = m.group(1).strip()
                            address = court_info[:m.start()].strip().rstrip(",").strip()
                        else:
                            court_name = court_info
                            address = ""

                        m_num = re.search(r'Судебный участок\s*(\d+)', court_name, flags=re.IGNORECASE)
                        dist_num = m_num.group(1) if m_num else ""

                        records.append({
                            "full_raw": court_info,
                            "address": address,
                            "court_name": court_name,
                            "district_num": dist_num,
                            "phone": phone,
                            "email": email
                        })
                if records:
                    break
            except Exception:
                pass

    return records


POSITIONS = [
    {"label": "Специалист 2 разряда", "code": "специалист 2р."},
    {"label": "Мировой судья", "code": "мировой судья"},
    {"label": "Помощник мирового судьи", "code": "помощник_судьи"},
    {"label": "Секретарь судебного заседания", "code": "секретарь"},
    {"label": "Заведующий канцелярией", "code": "завканц"},
    {"label": "Ввести должность вручную...", "code": "custom"}
]


def clean_for_filename(text: str) -> str:
    """Удаляет недопустимые для имен файлов Windows символы."""
    cleaned = re.sub(r'[\\/:*?"<>|]', '', text)
    return cleaned.strip()


def build_filename(district: str, role_code: str, address: str) -> str:
    """
    Формирует название файла по шаблону:
    акт_диагностики_2_участка_специалист 2р._г.Казань, ул.Авангардная, 143.docx
    """
    dist_str = district.strip()
    digits_match = re.search(r'(\d+)', dist_str)
    if digits_match:
        dist_part = f"{digits_match.group(1)}_участка"
    else:
        clean_dist = re.sub(r'[_ ]*участка', '', dist_str, flags=re.IGNORECASE).strip()
        dist_part = f"{clean_dist}_участка" if clean_dist else "участка"

    safe_role = clean_for_filename(role_code)
    safe_address = clean_for_filename(address)

    return f"акт_диагностики_{dist_part}_{safe_role}_{safe_address}.docx"


def apply_tnr_font(run, size_pt=14, bold=False):
    """Применяет шрифт Times New Roman 14 pt к фрагменту текста с поддержкой кириллицы."""
    run.font.name = "Times New Roman"
    run.font.size = Pt(size_pt)
    run.bold = bold

    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rFonts.set(qn("w:cs"), "Times New Roman")
    rFonts.set(qn("w:eastAsia"), "Times New Roman")


def create_docx(info_header: str, complaints: str, diagnostics: str, conclusion: str, images, include_labels=True):
    """
    Создает DOCX-документ со шрифтом Times New Roman 14:
    - Начальный информационный абзац: Акт диагностики по ...
    - Первый абзац: жалобы пользователя
    - Второй абзац: результаты диагностики
    - Блок фотографий (если прикреплены)
    - Последний абзац: вывод
    """
    doc = docx.Document()

    # Настройка базового стиля Normal
    normal_style = doc.styles["Normal"]
    if isinstance(normal_style, ParagraphStyle):
        normal_style.font.name = "Times New Roman"
        normal_style.font.size = Pt(14)
    rPr = normal_style.element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rFonts.set(qn("w:cs"), "Times New Roman")
    rFonts.set(qn("w:eastAsia"), "Times New Roman")

    # 0. Информационный вводный абзац в самом начале файла
    if info_header and info_header.strip():
        p_info = doc.add_paragraph()
        p_info.paragraph_format.space_after = Pt(12)
        p_info.paragraph_format.line_spacing = 1.15
        run_info = p_info.add_run(info_header.strip())
        apply_tnr_font(run_info, size_pt=14, bold=True)

    # 1. Первый абзац: жалобы пользователя
    p1 = doc.add_paragraph()
    p1.paragraph_format.space_after = Pt(8)
    p1.paragraph_format.line_spacing = 1.15
    if include_labels:
        run_lbl = p1.add_run("Жалобы пользователя: ")
        apply_tnr_font(run_lbl, size_pt=14, bold=True)
    run1 = p1.add_run(complaints.strip() if complaints else "Жалоб не поступало.")
    apply_tnr_font(run1, size_pt=14)

    # 2. Второй абзац: результаты диагностики
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(8)
    p2.paragraph_format.line_spacing = 1.15
    if include_labels:
        run_lbl = p2.add_run("Результаты диагностики: ")
        apply_tnr_font(run_lbl, size_pt=14, bold=True)
    run2 = p2.add_run(diagnostics.strip() if diagnostics else "Диагностика проведена.")
    apply_tnr_font(run2, size_pt=14)

    # 3. Фотографии (если прикреплены)
    if images:
        if include_labels:
            p_img_title = doc.add_paragraph()
            p_img_title.paragraph_format.space_after = Pt(6)
            run_img_lbl = p_img_title.add_run("Фотоматериалы:")
            apply_tnr_font(run_img_lbl, size_pt=14, bold=True)

        for idx, uploaded_img in enumerate(images, start=1):
            try:
                img_obj = Image.open(uploaded_img)
                img_io = io.BytesIO()
                if img_obj.mode in ("RGBA", "P"):
                    img_obj.save(img_io, format="PNG")
                else:
                    img_obj.save(img_io, format="JPEG", quality=92)
                img_io.seek(0)

                p_img = doc.add_paragraph()
                p_img.paragraph_format.space_after = Pt(8)
                p_img.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                doc.add_picture(img_io, width=Cm(15))

                if include_labels:
                    p_caption = doc.add_paragraph()
                    p_caption.paragraph_format.space_after = Pt(8)
                    p_caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run_cap = p_caption.add_run(f"Фото {idx}")
                    apply_tnr_font(run_cap, size_pt=14, bold=False)
            except Exception as e:
                p_err = doc.add_paragraph()
                run_err = p_err.add_run(f"[Ошибка добавления фото {idx}: {e}]")
                apply_tnr_font(run_err, size_pt=14)

    # 4. Последний абзац: вывод
    p_last = doc.add_paragraph()
    p_last.paragraph_format.space_after = Pt(8)
    p_last.paragraph_format.line_spacing = 1.15
    if include_labels:
        run_lbl = p_last.add_run("Вывод: ")
        apply_tnr_font(run_lbl, size_pt=14, bold=True)
    run_last = p_last.add_run(conclusion.strip() if conclusion else "Замечаний нет.")
    apply_tnr_font(run_last, size_pt=14)

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output


def diagnostic_act_page():
    try:
        st.set_page_config(
            page_title="Акт диагностики",
            page_icon="📝",
            layout="wide",
            initial_sidebar_state="expanded"
        )
    except Exception:
        pass

    # Принудительная светлая тема и стилизация элементов
    st.markdown("""
        <style>
            /* Светлая тема */
            .stApp {
                background-color: #FFFFFF;
                color: #1E293B;
            }
            .block-container {
                padding-top: 1.5rem;
                padding-bottom: 3rem;
            }
            .diag-card {
                background-color: #F8FAFC;
                border: 1px solid #E2E8F0;
                border-radius: 10px;
                padding: 16px;
                margin-bottom: 16px;
            }
            .info-box {
                background-color: #EFF6FF;
                border-left: 4px solid #3B82F6;
                padding: 12px 16px;
                border-radius: 6px;
                margin-bottom: 16px;
                font-size: 0.95rem;
            }
            /* Ограничение размера превью изображений на странице */
            [data-testid="stImage"] img {
                max-height: 160px !important;
                max-width: 200px !important;
                object-fit: cover !important;
                border-radius: 8px !important;
                border: 1px solid #CBD5E1 !important;
            }
            [data-testid="stImage"] {
                display: inline-block !important;
                margin-bottom: 8px !important;
            }
        </style>
    """, unsafe_allow_html=True)

    st.title("📝 Генератор акта диагностики оборудования")
    st.caption("Формирование акта диагностики с автоподстановкой адресов судебных участков РТ и экспортом в DOCX")

    # Загружаем базу участков
    courts_db = load_courts_database()

    # Список вариантов для выпадающего списка
    options = [
        f"{c['court_name']} — {c['address']} ({c.get('email', '')})"
        for c in courts_db
    ]

    # --- БЛОК 1: ВЫБОР УЧАСТКА И АВТОМАТИЧЕСКАЯ ПОДСТАНОВКА АДРЕСА ---
    st.subheader("1. Выбор судебного участка")

    col_select, col_search_hint = st.columns([3, 1])
    with col_select:
        selected_option = st.selectbox(
            "Судебный участок (выберите из базы участков РТ):",
            options=options,
            index=None,
            placeholder="Выберите судебный участок...",
            help="Начните вводить номер участка или район для быстрого поиска"
        )
    with col_search_hint:
        st.write("")
        st.write("")
        st.caption(f"🏛️ Всего в базе: **{len(courts_db)}** участков РТ")

    # Определяем выбранную запись
    chosen_idx = options.index(selected_option) if (selected_option and selected_option in options) else None
    selected_record = courts_db[chosen_idx] if (chosen_idx is not None and courts_db) else {}

    # Автоматически подставляемые значения в поля (с возможностью корректировки)
    col_dist, col_addr = st.columns([1, 2.5])
    with col_dist:
        district_val = st.text_input(
            "Номер участка:",
            value=selected_record.get("district_num", ""),
            placeholder="Например: 2",
            help="Номер участка (например: 2). В имени файла будет '2_участка'"
        )
    with col_addr:
        address_val = st.text_input(
            "Адрес судебного участка (подставляется автоматически):",
            value=selected_record.get("address", ""),
            placeholder="Адрес участка",
            help="Адрес участка, подставляемый из базы данных"
        )

    col_phone, col_email = st.columns(2)
    with col_phone:
        phone_val = st.text_input(
            "Телефон участка:",
            value=selected_record.get("phone", ""),
            placeholder="Телефон"
        )
    with col_email:
        email_val = st.text_input(
            "Электронная почта (email):",
            value=selected_record.get("email", ""),
            placeholder="email@tatar.ru"
        )

    court_full_name = selected_record.get("court_name", f"Судебный участок {district_val}" if district_val else "")

    # Формируем информационную строку
    # Формат: Акт диагностики по г.Казань, ул.Авангардная, 143, Судебный участок 2 по Приволжскому судебному району г.Казани, 8(843) 222-63-53, Ms.5002@tatar.ru
    info_parts = []
    if address_val.strip():
        info_parts.append(f"Акт диагностики по {address_val.strip()}")
    else:
        info_parts.append("Акт диагностики")
    if court_full_name.strip():
        info_parts.append(court_full_name.strip())
    if phone_val.strip():
        info_parts.append(phone_val.strip())
    if email_val.strip():
        info_parts.append(email_val.strip())

    formatted_info_header = ", ".join([p for p in info_parts if p])

    st.markdown(f"""
        <div class="info-box">
            📌 <strong>Информационный вводный абзац документа:</strong><br>
            <code>{formatted_info_header}</code>
        </div>
    """, unsafe_allow_html=True)

    st.markdown("---")

    # --- БЛОК 2: ДОЛЖНОСТЬ ---
    st.subheader("2. Должность сотрудника")
    pos_labels = [p["label"] for p in POSITIONS]
    selected_pos_label = st.selectbox("Выберите должность:", options=pos_labels, index=0)

    selected_pos_obj = next(p for p in POSITIONS if p["label"] == selected_pos_label)
    if selected_pos_obj["code"] == "custom":
        role_code = st.text_input("Введите должность для названия файла:", value="специалист 2р.")
    else:
        role_code = selected_pos_obj["code"]

    st.markdown("---")

    # --- БЛОК 3: ЖАЛОБЫ ПОЛЬЗОВАТЕЛЯ ---
    st.subheader("3. Жалобы пользователя (1 абзац)")
    complaints_text = st.text_area(
        "Опишите жалобы пользователя:",
        placeholder="Например: Не включается системный блок, шумит вентилятор охлаждения, медленно открываются документы в ПК МС...",
        height=120
    )

    # --- БЛОК 4: РЕЗУЛЬТАТЫ ДИАГНОСТИКИ ---
    st.subheader("4. Результаты диагностики (2 абзац)")
    diagnostics_text = st.text_area(
        "Опишите результаты проведенной диагностики:",
        placeholder="Например: В ходе диагностики выявлен выход из строя блока питания 400W (деформация конденсаторов). Жесткий диск и материнская плата исправны...",
        height=130
    )

    # --- БЛОК 5: ПРИКРЕПЛЕНИЕ ИЗОБРАЖЕНИЙ ---
    st.subheader("5. Прикрепление фотографий (или пропуск)")
    uploaded_images = st.file_uploader(
        "Прикрепите фотографии оборудования или экрана (необязательно, можно пропустить):",
        type=["jpg", "jpeg", "png", "webp", "bmp"],
        accept_multiple_files=True,
        help="Если прикрепление фото не требуется, пропустите этот блок."
    )

    if uploaded_images:
        st.success(f"Прикреплено фотографий: {len(uploaded_images)}")
        # Компактная галерея предпросмотра: фиксированная сетка из 4 колонок, чтобы фото не закрывали экран
        cols = st.columns(4)
        for i, img_file in enumerate(uploaded_images):
            with cols[i % 4]:
                st.image(img_file, caption=img_file.name, width=180)

    # --- БЛОК 6: ВЫВОД ---
    st.subheader("6. Вывод (последний абзац)")
    conclusion_text = st.text_area(
        "Опишите итоговый вывод и рекомендации:",
        placeholder="Например: Требуется замена блока питания на новый мощностью от 500W. Проведение профилактической чистки от пыли.",
        height=120
    )

    st.markdown("---")

    # --- БЛОК 7: ИМЯ ФАЙЛА И СКАЧИВАНИЕ ---
    st.subheader("7. Предварительный просмотр и скачивание DOCX")

    # Формируем имя файла
    filename = build_filename(district_val, role_code, address_val)

    st.info(f"📄 **Имя файла:** `{filename}`")

    col_opt1, col_opt2 = st.columns(2)
    with col_opt1:
        include_labels = st.checkbox(
            "Включать метки перед текстом («Жалобы пользователя:», «Результаты диагностики:», «Вывод:»)",
            value=True,
            help="При включении разделы оформляются жирным префиксом со шрифтом Times New Roman 14 pt"
        )

    # Кнопка формирования документа
    if st.button("📑 Сформировать акт (.docx)", type="primary"):
        if not complaints_text.strip() and not diagnostics_text.strip() and not conclusion_text.strip():
            st.warning("⚠️ Пожалуйста, заполните хотя бы одно из текстовых полей перед формированием акта.")
        else:
            with st.spinner("Формирование документа DOCX (Times New Roman 14)..."):
                try:
                    docx_stream = create_docx(
                        info_header=formatted_info_header,
                        complaints=complaints_text,
                        diagnostics=diagnostics_text,
                        conclusion=conclusion_text,
                        images=uploaded_images,
                        include_labels=include_labels
                    )
                    st.session_state["diag_docx_bytes"] = docx_stream.getvalue()
                    st.session_state["diag_docx_filename"] = filename
                    st.success("✅ Документ успешно сформирован!")
                except Exception as ex:
                    st.error(f"Ошибка формирования документа: {ex}")

    if "diag_docx_bytes" in st.session_state:
        st.download_button(
            label=f"💾 Скачать {st.session_state.get('diag_docx_filename', filename)}",
            data=st.session_state["diag_docx_bytes"],
            file_name=st.session_state.get("diag_docx_filename", filename),
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )


diagnostic_act_page()
